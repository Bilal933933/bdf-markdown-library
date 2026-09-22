import json
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pymupdf
import pytest
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.domains.conversion.models import ConversionStatus, ExtractionMethod
from app.domains.conversion.ocr import OCRResult
from app.domains.conversion.pipeline.process import (
    process_conversion,
    queue_conversion,
    requeue_conversion,
    requeue_orphans,
)
from app.domains.conversion.services.conversions import create_conversion
from app.infrastructure.database import (
    Base,
    ConversionRow,
    PageCheckpointRow,
    get_engine,
    get_session_factory,
)
from app.infrastructure.storage import LocalStorage

ARABIC_PARAGRAPH = (
    "المبتدأ اسم مرفوع يقع في أول الجملة والخبر يتمم معناه ويكمل الفائدة للمستمع والقارئ"
)


def build_arabic_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    page.insert_font(fontname="arb", fontfile="C:/Windows/Fonts/arialbd.ttf")
    page.insert_text((72, 72), "الوحدة الأولى", fontname="arb", fontsize=26)
    page.insert_text((72, 150), ARABIC_PARAGRAPH, fontname="arr", fontsize=12)
    second = doc.new_page()
    second.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    second.insert_text(
        (72, 72), "كان وأخواتها أفعال ناسخة ترفع المبتدأ وتنصب الخبر", fontname="arr", fontsize=12
    )
    return doc.tobytes()


def build_blank_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page()
    return doc.tobytes()


def build_scanned_pdf() -> bytes:
    src = pymupdf.open()
    page = src.new_page()
    page.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    page.insert_text((72, 72), ARABIC_PARAGRAPH, fontname="arr", fontsize=12)
    image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
    out = pymupdf.open()
    cover = out.new_page(width=page.rect.width, height=page.rect.height)
    cover.insert_image(cover.rect, stream=image)
    return out.tobytes()


def offline_settings() -> Settings:
    return Settings(gemini_keys="", gemini_api_key="")


def make_conversion(
    pdf: bytes, tmp_path: Path, filename: str = "book.pdf"
) -> tuple[str, Session, LocalStorage]:
    Base.metadata.create_all(get_engine())
    db = get_session_factory()()
    storage = LocalStorage(tmp_path)
    conversion = create_conversion(db, storage, filename, pdf, Settings())
    return conversion.id, db, storage


def cleanup(db: Session, conversion_id: str) -> None:
    try:
        db.query(PageCheckpointRow).filter_by(conversion_id=conversion_id).delete()
        db.query(ConversionRow).filter_by(id=conversion_id).delete()
        db.commit()
    finally:
        db.close()


def test_process_completes_and_writes_outputs(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        result = process_conversion(conversion_id, db, storage)
        assert result.status == ConversionStatus.COMPLETED
        assert result.progress == 100
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert "الوحدة الأولى" in markdown
        metadata = json.loads(storage.load(f"{conversion_id}/output/metadata.json"))
        assert metadata["page_count"] == 2
        assert metadata["stats"]["headings"] >= 1
        unit_meta = json.loads(storage.load(f"{conversion_id}/output/units/01/metadata.json"))
        assert unit_meta["title"] == "الوحدة الأولى"
        checkpoints = db.query(PageCheckpointRow).filter_by(conversion_id=conversion_id).all()
        assert [cp.status for cp in checkpoints] == ["done", "done"]
        assert all(cp.quality is not None for cp in checkpoints)
    finally:
        cleanup(db, conversion_id)


def test_completed_run_is_idempotent(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        first = process_conversion(conversion_id, db, storage)
        second = process_conversion(conversion_id, db, storage)
        assert (first.status, second.status) == (
            ConversionStatus.COMPLETED,
            ConversionStatus.COMPLETED,
        )
    finally:
        cleanup(db, conversion_id)


def test_blank_page_fails_conversion_without_network(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_blank_pdf(), tmp_path)
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.FAILED
        assert result.error is not None
        assert result.error.code == "PAGE_FAILED"
        requeued = requeue_conversion(db, conversion_id)
        assert requeued.status == ConversionStatus.QUEUED
        assert requeued.error is None
    finally:
        cleanup(db, conversion_id)


def test_requeue_orphans(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        row = db.get(ConversionRow, conversion_id)
        assert row is not None
        row.status = ConversionStatus.PROCESSING.value
        db.commit()
        assert requeue_orphans(db) == 1
        db.expire_all()
        assert db.get(ConversionRow, conversion_id).status == ConversionStatus.QUEUED.value
    finally:
        cleanup(db, conversion_id)


def test_fresh_heartbeat_is_not_requeued(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        row = db.get(ConversionRow, conversion_id)
        assert row is not None
        row.status = ConversionStatus.PROCESSING.value
        row.heartbeat_at = datetime.now(UTC)
        db.commit()
        assert requeue_orphans(db) == 0
        assert db.get(ConversionRow, conversion_id).status == ConversionStatus.PROCESSING.value
    finally:
        cleanup(db, conversion_id)


def test_stale_heartbeat_is_requeued(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        row = db.get(ConversionRow, conversion_id)
        assert row is not None
        row.status = ConversionStatus.PROCESSING.value
        row.heartbeat_at = datetime.now(UTC) - timedelta(seconds=301)
        db.commit()
        assert requeue_orphans(db, stale_after_seconds=300) == 1
    finally:
        cleanup(db, conversion_id)


def test_run_records_heartbeat(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_arabic_pdf(), tmp_path)
    try:
        process_conversion(conversion_id, db, storage, offline_settings())
        assert db.get(ConversionRow, conversion_id).heartbeat_at is not None
    finally:
        cleanup(db, conversion_id)


def test_mixed_pages_yield_partial_with_outputs(tmp_path: Path) -> None:
    doc = pymupdf.open()
    good = doc.new_page()
    good.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    good.insert_text((72, 72), ARABIC_PARAGRAPH, fontname="arr", fontsize=12)
    doc.new_page()
    conversion_id, db, storage = make_conversion(doc.tobytes(), tmp_path)
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.PARTIAL
        assert result.progress == 50
        assert result.error is not None and "2" in result.error.message
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert "المبتدأ" in markdown
        requeued = requeue_conversion(db, conversion_id)
        assert requeued.status == ConversionStatus.QUEUED
        again = process_conversion(conversion_id, db, storage, offline_settings())
        assert again.status == ConversionStatus.PARTIAL
    finally:
        cleanup(db, conversion_id)


def test_queue_returns_false_without_redis() -> None:
    assert queue_conversion("nope", Settings(redis_url=None)) is False


def build_docx_bytes() -> bytes:
    import io

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading("الوحدة الأولى", level=1)
    document.add_paragraph(
        "المبتدأ اسم مرفوع يقع في أول الجملة والخبر يتمم معناه ويكمل الفائدة للمستمع والقارئ"
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_completes_end_to_end(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_docx_bytes(), tmp_path, "doc.docx")
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.COMPLETED
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert markdown.startswith("# الوحدة الأولى")
    finally:
        cleanup(db, conversion_id)


def build_png_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    page.insert_text((72, 72), ARABIC_PARAGRAPH, fontname="arr", fontsize=12)
    return page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")


def test_png_completes_via_ocr(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.domains.conversion.pipeline.process.ocr_page_text", fake_ocr_page_text)
    conversion_id, db, storage = make_conversion(build_png_bytes(), tmp_path, "page.png")
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.COMPLETED
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert "المبتدأ" in markdown
    finally:
        cleanup(db, conversion_id)


def fake_ocr_page_text(
    image: bytes, mime: str, page_number: int, settings: Settings, **kwargs: object
) -> OCRResult:
    assert image.startswith(b"\xff\xd8")
    assert mime == "image/jpeg"
    return OCRResult(text=ARABIC_PARAGRAPH, confidence=None, method=ExtractionMethod.GEMINI)


def test_scanned_page_completes_via_ocr(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.domains.conversion.pipeline.process.ocr_page_text", fake_ocr_page_text)
    conversion_id, db, storage = make_conversion(build_scanned_pdf(), tmp_path)
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.COMPLETED
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert "المبتدأ" in markdown
        checkpoints = db.query(PageCheckpointRow).filter_by(conversion_id=conversion_id).all()
        assert [cp.method for cp in checkpoints] == [ExtractionMethod.GEMINI.value]
    finally:
        cleanup(db, conversion_id)


def test_webp_completes_via_ocr(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import io

    from PIL import Image

    monkeypatch.setattr("app.domains.conversion.pipeline.process.ocr_page_text", fake_ocr_page_text)
    buffer = io.BytesIO()
    Image.new("RGB", (60, 60), "white").save(buffer, format="WEBP")
    conversion_id, db, storage = make_conversion(buffer.getvalue(), tmp_path, "page.webp")
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.COMPLETED
    finally:
        cleanup(db, conversion_id)


def build_pdf_with_image() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="arr", fontfile="C:/Windows/Fonts/arial.ttf")
    page.insert_text((72, 72), ARABIC_PARAGRAPH, fontname="arr", fontsize=12)
    image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
    page.insert_image(pymupdf.Rect(72, 200, 300, 400), stream=image)
    return doc.tobytes()


def test_image_is_extracted_to_assets(tmp_path: Path) -> None:
    conversion_id, db, storage = make_conversion(build_pdf_with_image(), tmp_path)
    try:
        result = process_conversion(conversion_id, db, storage, offline_settings())
        assert result.status == ConversionStatus.COMPLETED
        manifest = json.loads(storage.load(f"{conversion_id}/output/assets.json"))
        assert len(manifest) == 1
        assert manifest[0]["id"] == "p1-img0"
        assert manifest[0]["mime"] == "image/png"
        asset_bytes = storage.load(f"{conversion_id}/output/assets/p1-img0.png")
        assert asset_bytes.startswith(b"\x89PNG")
        markdown = storage.load(f"{conversion_id}/output/document.md").decode("utf-8")
        assert f"![]({conversion_id}/output/assets/p1-img0.png)" in markdown
    finally:
        cleanup(db, conversion_id)
