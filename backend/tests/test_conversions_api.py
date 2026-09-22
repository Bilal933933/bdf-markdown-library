import io
from collections.abc import Iterator
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.api.v1.conversions import get_storage
from app.core.config import Settings
from app.core.errors.exceptions import ValidationError as AppValidationError
from app.domains.conversion.pipeline.process import QUEUE_NAME
from app.domains.conversion.services.conversions import create_conversion
from app.infrastructure.database import (
    Base,
    ConversionRow,
    PageCheckpointRow,
    get_engine,
    get_session_factory,
)
from app.infrastructure.database.session import get_db
from app.infrastructure.redis import get_redis, reset_redis
from app.infrastructure.storage import LocalStorage
from app.main import app


def purge_queue() -> None:
    """Drop jobs the API enqueued against a live server (dev hygiene)."""
    try:
        from rq import Queue

        Queue(QUEUE_NAME, connection=get_redis()).delete(delete_jobs=True)
    except Exception:
        pass
    finally:
        reset_redis()


def build_pdf(pages: int) -> bytes:
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page()
    return doc.tobytes()


@pytest.fixture()
def client(tmp_path: Path) -> Iterator[tuple[TestClient, list[str]]]:
    Base.metadata.create_all(get_engine())
    created: list[str] = []

    def override_db() -> Iterator[Session]:
        db = get_session_factory()()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_storage] = lambda: LocalStorage(tmp_path)
    try:
        yield TestClient(app), created
    finally:
        app.dependency_overrides.clear()
        purge_queue()
        db = get_session_factory()()
        try:
            for conversion_id in created:
                db.query(PageCheckpointRow).filter_by(conversion_id=conversion_id).delete()
                db.query(ConversionRow).filter_by(id=conversion_id).delete()
            db.commit()
        finally:
            db.close()


def test_post_pdf_returns_queued_with_checkpoints(client: tuple[TestClient, list[str]]) -> None:
    http, created = client
    response = http.post(
        "/api/v1/conversions", files={"file": ("book.pdf", build_pdf(2), "application/pdf")}
    )
    assert response.status_code == 202
    data = response.json()["data"]
    assert data["status"] == "queued"
    assert data["total_pages"] == 2
    assert response.json()["meta"]["request_id"] == response.headers["X-Request-ID"]
    created.append(data["id"])

    db = get_session_factory()()
    try:
        checkpoints = db.query(PageCheckpointRow).filter_by(conversion_id=data["id"]).all()
        assert [cp.status for cp in checkpoints] == ["pending", "pending"]
    finally:
        db.close()

    status = http.get(f"/api/v1/conversions/{data['id']}")
    assert status.status_code == 200
    assert status.json()["data"]["id"] == data["id"]


def test_get_unknown_conversion_returns_404(client: tuple[TestClient, list[str]]) -> None:
    http, _ = client
    response = http.get("/api/v1/conversions/does-not-exist")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "NOT_FOUND"


def test_post_unsupported_type_returns_422(client: tuple[TestClient, list[str]]) -> None:
    http, _ = client
    response = http.post(
        "/api/v1/conversions", files={"file": ("note.zip", b"PK\x05\x06data", "application/zip")}
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


def test_post_png_is_accepted(client: tuple[TestClient, list[str]]) -> None:
    http, created = client
    doc = pymupdf.open()
    doc.new_page()
    png = doc[0].get_pixmap().tobytes("png")
    response = http.post("/api/v1/conversions", files={"file": ("page.png", png, "image/png")})
    assert response.status_code == 202
    assert response.json()["data"]["total_pages"] == 1
    created.append(response.json()["data"]["id"])


def test_post_corrupt_image_returns_422(client: tuple[TestClient, list[str]]) -> None:
    http, _ = client
    response = http.post(
        "/api/v1/conversions", files={"file": ("page.png", b"not-an-image", "image/png")}
    )
    assert response.status_code == 422


def test_post_webp_is_accepted(client: tuple[TestClient, list[str]]) -> None:
    from PIL import Image

    http, created = client
    buffer = io.BytesIO()
    Image.new("RGB", (60, 60), "white").save(buffer, format="WEBP")
    response = http.post(
        "/api/v1/conversions", files={"file": ("page.webp", buffer.getvalue(), "image/webp")}
    )
    assert response.status_code == 202
    assert response.json()["data"]["total_pages"] == 1
    created.append(response.json()["data"]["id"])


def test_post_txt_is_accepted(client: tuple[TestClient, list[str]]) -> None:
    http, created = client
    response = http.post(
        "/api/v1/conversions",
        files={"file": ("note.txt", "سطر أول.\n\nسطر ثان.".encode(), "text/plain")},
    )
    assert response.status_code == 202
    data = response.json()["data"]
    assert data["total_pages"] == 1
    created.append(data["id"])


def test_post_docx_is_accepted(client: tuple[TestClient, list[str]]) -> None:
    from docx import Document as DocxDocument

    http, created = client
    document = DocxDocument()
    document.add_heading("عنوان", level=1)
    document.add_paragraph("فقرة تجريبية للتحويل.")
    buffer = io.BytesIO()
    document.save(buffer)
    response = http.post(
        "/api/v1/conversions",
        files={
            "file": (
                "doc.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 202
    assert response.json()["data"]["total_pages"] == 1
    created.append(response.json()["data"]["id"])


def test_service_rejects_oversize_and_over_pages(tmp_path: Path) -> None:
    Base.metadata.create_all(get_engine())
    db = get_session_factory()()
    try:
        with pytest.raises(AppValidationError, match="size limit"):
            create_conversion(
                db,
                LocalStorage(tmp_path),
                "book.pdf",
                build_pdf(1),
                Settings(max_file_size=10),
            )
        db.rollback()
        with pytest.raises(AppValidationError, match="page limit"):
            create_conversion(
                db, LocalStorage(tmp_path), "book.pdf", build_pdf(2), Settings(max_pages=1)
            )
        db.rollback()
    finally:
        db.close()


def post_pdf(http: TestClient, created: list[str], name: str = "book.pdf") -> str:
    response = http.post("/api/v1/conversions", files={"file": (name, build_pdf(1))})
    assert response.status_code == 202
    conversion_id = response.json()["data"]["id"]
    created.append(conversion_id)
    return conversion_id


def test_list_returns_newest_first_with_pagination(
    client: tuple[TestClient, list[str]],
) -> None:
    http, created = client
    first = post_pdf(http, created, "a.pdf")
    second = post_pdf(http, created, "b.pdf")
    third = post_pdf(http, created, "c.pdf")

    listed = http.get("/api/v1/conversions").json()["data"]
    ids = [item["id"] for item in listed]
    assert ids.index(third) < ids.index(second) < ids.index(first)

    page = http.get("/api/v1/conversions", params={"limit": 2}).json()["data"]
    assert len(page) == 2
    rest = http.get("/api/v1/conversions", params={"limit": 2, "offset": 2}).json()["data"]
    assert {item["id"] for item in page} | {item["id"] for item in rest} >= {first, second, third}

    assert http.get("/api/v1/conversions", params={"limit": 0}).status_code == 422
    assert http.get("/api/v1/conversions", params={"limit": 101}).status_code == 422


def test_cors_allows_configured_origin(monkeypatch: pytest.MonkeyPatch) -> None:
    import app.main as main_module

    monkeypatch.setattr(
        main_module, "get_settings", lambda: Settings(cors_origins="https://x.example")
    )
    client = TestClient(main_module.create_app())
    response = client.options(
        "/api/v1/conversions",
        headers={
            "Origin": "https://x.example",
            "Access-Control-Request-Method": "POST",
        },
    )
    assert response.headers["access-control-allow-origin"] == "https://x.example"


def test_cors_absent_by_default() -> None:
    client = TestClient(app)
    response = client.options(
        "/api/v1/conversions",
        headers={"Origin": "https://x.example", "Access-Control-Request-Method": "POST"},
    )
    assert "access-control-allow-origin" not in response.headers
