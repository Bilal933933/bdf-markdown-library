import io

import pytest

from app.core.errors.exceptions import ValidationError as AppValidationError
from app.domains.conversion.extract import decode_text, extract_docx, extract_text
from app.domains.conversion.models import BlockType


def build_docx() -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading("الوحدة", level=1)
    document.add_heading("درس", level=2)
    document.add_paragraph("فقرة أولى.")
    document.add_paragraph("نقطة", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "مصطلح"
    table.cell(0, 1).text = "تعريف"
    table.cell(1, 0).text = "المبتدأ"
    table.cell(1, 1).text = "اسم مرفوع"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_extracts_headings_list_table_in_order() -> None:
    blocks = extract_docx(build_docx(), "doc.docx")
    kinds = [(block.type, getattr(block.payload, "text", "")) for block in blocks]
    assert kinds[0] == (BlockType.HEADING, "الوحدة")
    assert blocks[0].payload.level == 1
    assert kinds[1][0] == BlockType.HEADING
    assert kinds[2] == (BlockType.PARAGRAPH, "فقرة أولى.")
    assert blocks[3].type == BlockType.LIST
    assert blocks[3].payload.items[0].text == "نقطة"
    assert blocks[4].type == BlockType.TABLE
    assert blocks[4].payload.rows[1].cells[0].text == "المبتدأ"
    assert [block.order for block in blocks] == list(range(len(blocks)))
    assert all(block.source.method.value == "docx" for block in blocks)


def test_docx_rejects_garbage() -> None:
    with pytest.raises(AppValidationError, match="Invalid DOCX"):
        extract_docx(b"PK\x03\x04not-a-docx", "doc.docx")


def test_text_splits_and_decodes() -> None:
    blocks = extract_text("سطر أول.\n\nسطر ثان.".encode(), "note.txt")
    assert [block.payload.text for block in blocks] == ["سطر أول.", "سطر ثان."]
    assert decode_text("نص عربي.".encode("cp1256")) == "نص عربي."
