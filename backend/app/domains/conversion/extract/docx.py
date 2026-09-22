"""DOCX extraction — headings, lists, tables, paragraphs to Blocks (pure).

Body order follows the document XML sequence (paragraphs and tables interleaved).
"""

import io
from typing import Any

from app.core.errors.exceptions import ValidationError
from app.domains.conversion.models import (
    Block,
    BlockSource,
    BlockType,
    Cell,
    ExtractionMethod,
    HeadingPayload,
    ListItem,
    ListPayload,
    ParagraphPayload,
    Row,
    TablePayload,
)


def _source(filename: str) -> BlockSource:
    return BlockSource(file=filename, pages=[1], method=ExtractionMethod.DOCX)


def _heading_level(style_name: str) -> int:
    for token in style_name.split():
        if token.isdigit():
            return min(max(int(token), 1), 6)
    return 1


def extract_docx(data: bytes, filename: str) -> list[Block]:
    try:
        from docx import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph as DocxParagraph
    except ImportError as exc:
        raise ValidationError("python-docx is not installed") from exc
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError("Invalid DOCX file", details={"filename": filename}) from exc

    blocks: list[Block] = []

    def add(payload: Any, block_type: BlockType) -> None:
        blocks.append(
            Block(
                id=f"p1-b{len(blocks)}",
                type=block_type,
                order=len(blocks),
                source=_source(filename),
                payload=payload,
            )
        )

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            _add_paragraph(DocxParagraph(child, document), add)
        elif isinstance(child, CT_Tbl):
            table = DocxTable(child, document)
            rows = [
                Row(cells=[Cell(text=cell.text.strip()) for cell in row.cells])
                for row in table.rows
            ]
            rows = [row for row in rows if any(cell.text for cell in row.cells)]
            if rows:
                add(TablePayload(rows=rows), BlockType.TABLE)
    return blocks


def _add_paragraph(paragraph: Any, add: Any) -> None:
    text = paragraph.text.strip()
    if not text:
        return
    style = paragraph.style.name if paragraph.style else ""
    if style.startswith("Heading"):
        add(HeadingPayload(level=_heading_level(style), text=text), BlockType.HEADING)
    elif "List" in style:
        ordered = "Number" in style
        add(
            ListPayload(ordered=ordered, items=[ListItem(text=text)]),
            BlockType.LIST,
        )
    else:
        add(ParagraphPayload(text=text), BlockType.PARAGRAPH)
